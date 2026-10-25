"""Exports: formula-injection escaping, file names, the export job, the filesystem store and production gating."""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

import pytest
from authkit import firm_db
from conftest import START, TENANT_A, TEST_KEYS_JSON, TEST_KID, FakeClock
from datakit import CARRIER, created, error_code, seed_firms
from fastapi import FastAPI
from fastapi.testclient import TestClient
from importkit import fresh_owner, run_jobs, upload
from openpyxl import load_workbook
from pydantic import SecretStr

from manceps_api.app import create_app
from manceps_api.exports import writer
from manceps_api.exports.store import ExportStore, FilesystemExportStore, StoreError, store_from_settings
from manceps_api.integrations.configured import Integrations
from manceps_api.settings import Settings
from manceps_core.mongo import MongoClientT

FORMULA_NAME = "=cmd|' /C calc'!A0"


@pytest.fixture
def clock() -> FakeClock:
    """A fake clock starting at the real current time (whole seconds).

    Export rows carry a TTL on `expires_at`, and mongomock (like MongoDB) expires them against the real clock, so
    these tests must not start from the fixed date the rest of the suite uses.

    Returns:
        The clock.
    """
    return FakeClock(datetime.now(UTC).replace(microsecond=0))


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    """Test settings with a temporary export directory.

    Args:
        tmp_path: Pytest temporary directory.

    Returns:
        The settings.
    """
    return Settings(env="test", jwt_active_kid=TEST_KID, jwt_keys_json=SecretStr(TEST_KEYS_JSON),
                    export_dir=str(tmp_path / "exports"))


def _store(app: FastAPI) -> ExportStore:
    """Return the app's export store.

    Args:
        app: The app.

    Returns:
        The store.
    """
    store = app.state.context.export_store
    assert store is not None
    return store


def _export(client: TestClient, headers: dict[str, str], body: dict[str, Any]) -> dict[str, Any]:
    """POST an export and expect 202.

    Args:
        client: Test client.
        headers: Auth headers.
        body: Request body.

    Returns:
        The export.
    """
    response = client.post("/v1/exports", json=body, headers=headers)
    assert response.status_code == 202, response.text
    return response.json()


@pytest.mark.parametrize(("value", "escaped"), [
    ("=1+1", "'=1+1"), ("+12145550100", "'+12145550100"), ("-5", "'-5"), ("@SUM(A1)", "'@SUM(A1)"),
    ("\tx", "'\tx"), ("\rx", "'\rx"), ("safe", "safe"), ("", ""), (5, 5), (None, None), (True, True),
])
def test_formula_injection_escaped(value: Any, escaped: Any) -> None:
    assert writer.escape_cell(value) == escaped


def test_formula_injection_escaped_in_exported_files(client: TestClient, mongo_client: MongoClientT,
                                                     clock: FakeClock, app: FastAPI) -> None:
    firms = seed_firms(client, mongo_client, clock)
    owner = firms.headers("owner")
    created(client, owner, "/v1/carriers", {**CARRIER, "name": FORMULA_NAME, "notes": "@risky"})
    csv_export = _export(client, owner, {"entity": "carriers", "format": "csv"})
    xlsx_export = _export(client, owner, {"entity": "carriers", "format": "xlsx"})
    assert run_jobs(mongo_client, clock, _store(app)) == 2
    text = client.get(f"/v1/exports/{csv_export['id']}/download", headers=owner).text
    rows = list(csv.DictReader(io.StringIO(text)))
    assert rows[0]["name"] == "'" + FORMULA_NAME and rows[0]["notes"] == "'@risky"
    assert rows[0]["phone"] == "'+12145550300"
    content = client.get(f"/v1/exports/{xlsx_export['id']}/download", headers=owner).content
    sheet = load_workbook(io.BytesIO(content), read_only=False).active
    assert sheet is not None
    header = [cell.value for cell in sheet[1]]
    name_cell = sheet.cell(row=2, column=header.index("name") + 1)
    assert name_cell.value == "'" + FORMULA_NAME and name_cell.data_type == "s"


def test_export_filename_is_sanitized() -> None:
    assert writer.export_filename("carriers", "csv", START) == "carriers-20501025-120000.csv"
    assert writer.content_disposition("bad\"name\r\n.csv") == 'attachment; filename="bad_name_.csv"'


def test_export_flow_with_filters(client: TestClient, mongo_client: MongoClientT, clock: FakeClock,
                                  app: FastAPI) -> None:
    firms = seed_firms(client, mongo_client, clock)
    owner = firms.headers("owner")
    created(client, owner, "/v1/carriers", CARRIER)
    other = created(client, owner, "/v1/carriers", {"phone": "2145550399", "name": "Paused One"})
    assert client.patch(f"/v1/carriers/{other['id']}", json={"status": "paused"}, headers=owner).status_code == 200
    created_at = clock()
    queued = _export(client, firms.headers("dispatcher"), {"entity": "carriers", "format": "csv",
                                                           "filters": {"status": "paused"}})
    assert queued["status"] == "queued" and queued["download_url"] is None
    early = client.get(f"/v1/exports/{queued['id']}/download", headers=owner)
    assert early.status_code == 409 and error_code(early) == "export_not_ready"
    run_jobs(mongo_client, clock, _store(app))
    ready = client.get(f"/v1/exports/{queued['id']}", headers=firms.headers("viewer")).json()
    assert ready["status"] == "ready" and ready["rows"] == 1
    assert ready["download_url"] == f"/v1/exports/{queued['id']}/download"
    download = client.get(ready["download_url"], headers=firms.headers("viewer"))
    assert download.status_code == 200 and download.headers["content-type"].startswith("text/csv")
    expected = writer.export_filename("carriers", "csv", created_at)
    assert f'filename="{expected}"' in download.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(download.text)))
    assert [row["name"] for row in rows] == ["Paused One"]
    clock.advance(timedelta(days=2))
    expired = client.get(ready["download_url"], headers=fresh_owner(client))
    assert expired.status_code == 410 and error_code(expired) == "export_expired"


def test_export_every_entity(client: TestClient, mongo_client: MongoClientT, clock: FakeClock, app: FastAPI) -> None:
    firms = seed_firms(client, mongo_client, clock)
    owner = firms.headers("owner")
    carrier = created(client, owner, "/v1/carriers", CARRIER)
    created(client, owner, "/v1/drivers", {"carrier_id": carrier["id"], "name": "Sam", "phone": "2145550900"})
    created(client, owner, "/v1/trucks", {"carrier_id": carrier["id"], "unit": "T1"})
    created(client, owner, "/v1/brokers", {"mc_number": "123456", "company": "Acme"})
    created(client, owner, "/v1/contacts", {"kind": "shipper", "name": "Dock", "phones": ["2145550901"]})
    created(client, owner, "/v1/facilities", {"name": "DC 1", "state": "TX"})
    ids = [_export(client, owner, {"entity": entity, "format": "xlsx"})["id"]
           for entity in ("carriers", "drivers", "trucks", "brokers", "contacts", "facilities")]
    assert run_jobs(mongo_client, clock, _store(app)) == len(ids)
    for export_id in ids:
        body = client.get(f"/v1/exports/{export_id}", headers=owner).json()
        assert body["status"] == "ready" and body["rows"] == 1, body


def test_production_without_store_is_503(mongo_client: MongoClientT, clock: FakeClock,
                                         integrations: Integrations) -> None:
    production = Settings(env="production", jwt_active_kid=TEST_KID, jwt_keys_json=SecretStr(TEST_KEYS_JSON),
                          mongodb_uri=SecretStr("mongodb://unused"), secrets_key=SecretStr("k" * 43),
                          public_base_url="https://api.example")
    assert store_from_settings(production) is None
    app = create_app(production, mongo_client, clock=clock, integrations=integrations)
    with TestClient(app) as client:
        firms = seed_firms(client, mongo_client, clock)
        response = client.post("/v1/exports", json={"entity": "carriers", "format": "csv"},
                               headers=firms.headers("owner"))
        assert response.status_code == 503 and error_code(response) == "exports_not_configured"


def test_export_idempotency_key_replays(client: TestClient, mongo_client: MongoClientT, clock: FakeClock) -> None:
    firms = seed_firms(client, mongo_client, clock)
    headers = {**firms.headers("owner"), "Idempotency-Key": "export-key-0001"}
    first = client.post("/v1/exports", json={"entity": "brokers", "format": "csv"}, headers=headers)
    again = client.post("/v1/exports", json={"entity": "brokers", "format": "csv"}, headers=headers)
    assert first.status_code == again.status_code == 202
    assert again.headers["Idempotency-Replayed"] == "true" and again.json()["id"] == first.json()["id"]
    assert firm_db(mongo_client, TENANT_A)["exports"].count_documents({}) == 1


def test_filesystem_store_confines_keys(tmp_path: Path) -> None:
    store = FilesystemExportStore(tmp_path / "root")
    source = tmp_path / "file.csv"
    source.write_bytes(b"a,b\n")
    key = "apex-freight/0123456789abcdef01234567/carriers.csv"
    assert store.put(key, source) == 4
    with store.open(key) as handle:
        assert handle.read() == b"a,b\n"
    assert store.url_for(key, timedelta(minutes=5)) is None
    for bad in ("../x/0123456789abcdef01234567/a.csv", "apex-freight/../../etc/passwd", "/abs/path",
                "apex-freight/0123456789abcdef01234567/../../x.csv", ""):
        with pytest.raises(StoreError):
            store.put(bad, source)
    with pytest.raises(StoreError):
        store.open("apex-freight/0123456789abcdef01234567/missing.csv")
    store.delete(key)
    store.delete(key)
    assert not (tmp_path / "root" / key).exists()


def test_failed_export_is_recorded(client: TestClient, mongo_client: MongoClientT, clock: FakeClock,
                                   tmp_path: Path) -> None:
    firms = seed_firms(client, mongo_client, clock)
    owner = firms.headers("owner")
    queued = _export(client, owner, {"entity": "carriers", "format": "csv"})
    blocker = tmp_path / "not-a-dir"
    blocker.write_text("x")
    run_jobs(mongo_client, clock, FilesystemExportStore(blocker))
    body = client.get(f"/v1/exports/{queued['id']}", headers=owner).json()
    assert body["status"] == "failed" and body["error"] == "store_failed"


def test_import_source_filename_is_sanitized(client: TestClient, mongo_client: MongoClientT, clock: FakeClock) -> None:
    firms = seed_firms(client, mongo_client, clock)
    content = b"Name,Phone\na,2145550100\n"
    response = upload(client, firms.headers("owner"), "carriers", "../../evil name;.csv", content)
    assert response.status_code == 201 and response.json()["source_filename"] == "evil_name.csv"
